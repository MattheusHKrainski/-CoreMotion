# -*- coding: utf-8 -*-
"""
build_deck.py — Gera a apresentação acadêmica do projeto CoreMotiom (.pptx).

Uso:  .venv/bin/python docs/apresentacao/build_deck.py
Tema: claro/corporativo, 16:9, 11 slides.
Todos os números citados foram medidos no repositório (outubro/2026).
"""
import os

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

BASE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(os.path.dirname(BASE))
IMG = os.path.join(REPO, "public", "products", "nike-air-force-1.jpg")
OUT = os.path.join(BASE, "CoreMotiom-Apresentacao-Academica.pptx")

# ---- paleta (claro corporativo) ----
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x0F, 0x17, 0x2A)
BODY = RGBColor(0x33, 0x41, 0x55)
MUTED = RGBColor(0x64, 0x74, 0x8B)
RED = RGBColor(0xDC, 0x26, 0x26)
RED_SOFT = RGBColor(0xFE, 0xF2, 0xF2)
CARD = RGBColor(0xF1, 0xF5, 0xF9)
LINE = RGBColor(0xE2, 0xE8, 0xF0)

FONT = "Calibri"
W, H = 13.333, 7.5
N_SLIDES = 11

prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
BLANK = prs.slide_layouts[6]

_counter = {"n": 0}


def slide():
    _counter["n"] += 1
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(W), Inches(H))
    bg.fill.solid()
    bg.fill.fore_color.rgb = WHITE
    bg.line.fill.background()
    bg.shadow.inherit = False
    return s


def rect(s, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, radius=None):
    sp = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(1)
    sp.shadow.inherit = False
    return sp


def text(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         space_after=0, line_spacing=1.0):
    """runs: lista de parágrafos; cada parágrafo = lista de (texto, size, color, bold)."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        p.space_before = Pt(0)
        p.line_spacing = line_spacing
        for (t, size, color, bold) in para:
            r = p.add_run()
            r.text = t
            r.font.size = Pt(size)
            r.font.color.rgb = color
            r.font.bold = bold
            r.font.name = FONT
    return tb


def header(s, kicker, title):
    rect(s, 0.6, 0.52, 0.28, 0.28, fill=RED)
    text(s, 1.0, 0.5, 11.5, 0.32,
         [[(kicker.upper(), 11, RED, True)]])
    text(s, 0.6, 0.82, 12.1, 0.7, [[(title, 29, INK, True)]])
    rect(s, 0.62, 1.52, 1.15, 0.05, fill=RED)


def footer(s):
    n = _counter["n"]
    text(s, 0.6, 7.05, 9.0, 0.3,
         [[("CoreMotiom — Apresentação Acadêmica de Projeto  •  Curitiba, outubro de 2026", 9.5, MUTED, False)]])
    text(s, 11.7, 7.05, 1.05, 0.3, [[(f"{n} / {N_SLIDES}", 9.5, MUTED, False)]],
         align=PP_ALIGN.RIGHT)


def bullets(s, x, y, w, h, items, size=15, color=BODY, after=9, lh=1.05):
    paras = []
    for it in items:
        paras.append([("▪  ", 13, RED, True), (it, size, color, False)])
    text(s, x, y, w, h, paras, space_after=after, line_spacing=lh)


def card(s, x, y, w, h, title, body, title_size=14, body_size=12, fill=CARD,
         line=LINE, body_color=BODY, top_accent=True):
    rect(s, x, y, w, h, fill=fill, line=line, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    if top_accent:
        rect(s, x + 0.18, y + 0.16, 0.22, 0.22, fill=RED)
    text(s, x + 0.5, y + 0.13, w - 0.7, 0.4, [[(title, title_size, INK, True)]])
    if body:
        text(s, x + 0.22, y + 0.52, w - 0.44, h - 0.62,
             [[(line_t, body_size, body_color, False)] for line_t in body],
             space_after=4, line_spacing=1.02)


# ============================================================= S1 — capa
s = slide()
rect(s, 0, 0, W, 0.18, fill=RED)
rect(s, 8.15, 1.15, 4.6, 5.4, fill=CARD, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
if os.path.exists(IMG):
    s.shapes.add_picture(IMG, Inches(8.45), Inches(1.55), Inches(4.0), Inches(4.0))
text(s, 8.45, 5.7, 4.0, 0.8,
     [[("Catálogo do marketplace:", 10.5, MUTED, False)],
      [("produtos com fotos reais no repositório", 10.5, MUTED, False)]],
     align=PP_ALIGN.CENTER)

text(s, 0.75, 1.5, 7.2, 0.35, [[("PROJETO ACADÊMICO • ENGENHARIA DE SOFTWARE", 12, RED, True)]])
text(s, 0.72, 1.95, 7.3, 1.1, [[("CoreMotiom", 48, INK, True)]])
text(s, 0.75, 3.05, 7.1, 0.9,
     [[("Marketplace & Plataforma Esportiva", 20, RED, True)],
      [("de Alta Performance", 20, RED, True)]])
rect(s, 0.78, 4.15, 1.4, 0.05, fill=RED)
text(s, 0.75, 4.4, 7.1, 1.2,
     [[("Lojas oficiais verificadas · venda C2C entre atletas ·", 13.5, BODY, False)],
      [("custódia financeira (escrow) · treinadores e comunidade", 13.5, BODY, False)]],
     line_spacing=1.15)
text(s, 0.75, 5.9, 7.1, 0.7,
     [[("Desenvolvimento: ", 12.5, MUTED, False), ("Mattheus", 12.5, INK, True),
       ("  ·  github.com/MattheusHKrainski", 12.5, MUTED, False)],
      [("Curitiba — outubro de 2026", 12.5, MUTED, False)]],
     line_spacing=1.3)
# capa é o slide 1 (sem rodapé); as demais seguem 2..11

# ============================================================= S2 — problema
s = slide()
header(s, "Contextualização", "O problema que o projeto resolve")
bullets(s, 0.6, 1.85, 6.5, 4.6, [
    "Equipamentos de alta performance (tênis com placa de carbono, GPS, compressão) têm alto custo e forte mercado de revenda.",
    "A revenda entre atletas ocorre em canais genéricos: sem inspeção de desgaste, sem preço de referência e com risco de golpe.",
    "Lojas esportivas online convivem com falsificações e não possuem selo de verificação confiável no varejo generalista.",
    "Não havia plataforma vertical integrando marketplace, custódia financeira, treinadores e comunidade em um só produto.",
], size=14.5)
card(s, 7.5, 1.85, 5.25, 4.35, "Evidências no código", [
    "Tela “Vender C2C” bloqueia visitantes:",
    "“proteção anti-golpe” (app/page.tsx)",
    "",
    "Solicitação de selo oficial com CNPJ",
    "(components/StoresView.tsx)",
    "",
    "Status de pedido com escrow:",
    "escrow_locked (types/supabase.ts)",
], body_size=12)
footer(s)

# ============================================================= S3 — objetivos
s = slide()
header(s, "Objetivos", "Objetivo geral e específicos")
card(s, 0.6, 1.85, 12.13, 1.15, "Objetivo geral", [
    "Desenvolver uma plataforma web que integre um marketplace esportivo híbrido (B2C + C2C) com custódia financeira, "
    "lojas oficiais verificadas e serviços para atletas.",
], title_size=14, body_size=13.5)
text(s, 0.62, 3.3, 12.0, 0.4, [[("OBJETIVOS ESPECÍFICOS", 11.5, RED, True)]])
bullets(s, 0.6, 3.72, 12.1, 3.0, [
    "Implementar verificação de lojas oficiais com CNPJ e selo “Verificado”.",
    "Proteger vendas C2C com pagamento retido em custódia (escrow) até a confirmação do pedido.",
    "Oferecer checkout com PIX (5% de desconto), cartão de crédito e boleto.",
    "Integrar comunidade esportiva, treinadores com CREF e notícias ao ecossistema.",
    "Garantir segurança de dados com Row Level Security (RLS) e perfis de acesso.",
    "Prover painel administrativo de auditoria, moderação e custódia.",
], size=14, after=7)
footer(s)

# ============================================================= S4 — justificativa
s = slide()
header(s, "Justificativa", "Por que este projeto é relevante")
bullets(s, 0.6, 1.85, 12.1, 4.8, [
    "O crescimento da corrida de rua e dos esportes de performance no Brasil amplia o mercado de equipamentos novos e seminovos.",
    "Confiança é o principal entrave da revenda C2C: assimetria de informação sobre desgaste, autenticidade e preço justo.",
    "Supabase com RLS permite segurança em nível de banco de dados — prática de produção aplicada já no MVP acadêmico.",
    "O projeto consolida competências de engenharia: React/Next.js, TypeScript estrito, SQL, segurança e UX.",
    "Serve como base replicável: governança de Git, migrações versionadas e tipagem de ponta a ponta documentadas em CONTRIBUTING.md.",
], size=15.5, after=12)
footer(s)

# ============================================================= S5 — metodologia
s = slide()
header(s, "Metodologia", "Processo de desenvolvimento e qualidade")
bullets(s, 0.6, 1.85, 12.1, 4.8, [
    "Git Flow pragmático (equipe de 5): branch main protegida; merge apenas via Pull Request com revisão e verificação (CONTRIBUTING.md).",
    "TypeScript em modo estrito (strict: true), sem any injustificado; tipagem do banco gerada automaticamente (types/supabase.ts).",
    "Validação de ponta a ponta com schemas Zod + React Hook Form, centralizados em lib/schemas/.",
    "Banco de dados governado: migrações versionadas em supabase/migrations/ e RLS obrigatório em toda tabela nova.",
    "Camada de serviços isolada da UI (services/*.ts), permitindo trocar Supabase/fallback local sem reescrever telas.",
    "Pipeline de qualidade acionável localmente: npm run ci = lint + type-check + build.",
], size=14.5, after=10)
footer(s)

# ============================================================= S6 — stack
s = slide()
header(s, "Fundamentação técnica", "Stack tecnológica")
card(s, 0.6, 1.85, 6.0, 2.3, "Front-end", [
    "Next.js 15 — App Router, Server Components e Server Actions",
    "React 19 · Tailwind CSS v4 · shadcn/ui (Radix UI)",
    "motion (animações) · lucide-react (ícones)",
])
card(s, 6.73, 1.85, 6.0, 2.3, "Back-end / BaaS", [
    "Supabase: PostgreSQL 15+, Auth, Storage e RLS",
    "Server Actions (app/actions/) e API Routes (app/api/)",
    "node-postgres (pg) em scripts de seed e migração",
])
card(s, 0.6, 4.35, 6.0, 2.3, "Qualidade & validação", [
    "TypeScript estrito · ESLint (eslint-config-next)",
    "Zod (schemas) · React Hook Form (+ @hookform/resolvers)",
    "type-check e build integrados no script ci",
])
card(s, 6.73, 4.35, 6.0, 2.3, "Inteligência artificial", [
    "@google/genai — API Gemini (dependência declarada)",
    "SmartScan: avaliação de desgaste e preço justo",
    "(componente prototipado, em fase de integração)",
])
footer(s)

# ============================================================= S7 — arquitetura
s = slide()
header(s, "Arquitetura", "Visão em camadas do sistema")
rows = [
    ("UI / Views", "Marketplace · Lojas · Vender C2C · Coaches · Comunidade · Admin (components/)"),
    ("Estado global", "CoreMotiomProvider (React Context): sessão, carrinho, toasts, modais — lib/store.tsx"),
    ("Integração", "services/*.ts · Server Actions (app/actions/) · API Routes (app/api/)"),
    ("Dados", "Supabase: Auth · PostgreSQL + RLS · Storage  |  fallback local (lib/initial-data.ts)"),
]
y = 1.9
for i, (t, b) in enumerate(rows):
    fill = RED_SOFT if i % 2 else CARD
    rect(s, 1.1, y, 10.0, 0.92, fill=fill, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, 1.35, y + 0.1, 2.3, 0.72, [[(t, 13.5, RED if i == 3 else INK, True)]], anchor=MSO_ANCHOR.MIDDLE)
    text(s, 3.7, y + 0.1, 7.2, 0.72, [[(b, 12, BODY, False)]], anchor=MSO_ANCHOR.MIDDLE)
    if i < len(rows) - 1:
        text(s, 5.9, y + 0.86, 0.6, 0.4, [[("▼", 12, RED, True)]], align=PP_ALIGN.CENTER)
    y += 1.22
rect(s, 11.35, 2.6, 1.5, 2.2, fill=WHITE, line=RED, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 11.35, 2.9, 1.5, 1.6,
     [[("Gemini", 13, RED, True)], [("API", 13, RED, True)], [("SmartScan", 10.5, MUTED, False)], [("(IA)", 10.5, MUTED, False)]],
     align=PP_ALIGN.CENTER, line_spacing=1.05)
footer(s)

# ============================================================= S8 — dados & segurança
s = slide()
header(s, "Banco de dados & segurança", "Modelo relacional e RLS")
tables = [
    ("profiles", "papéis: admin · merchant · seller · user; sincronizada com auth.users"),
    ("stores", "verificação: none → pending → verified / rejected; docs + CNPJ"),
    ("products", "product_type: b2c | c2c · condição · estoque · status active/sold/suspended"),
    ("orders", "escrow: pending_payment → escrow_locked → preparing → shipped → delivered"),
    ("community_posts", "categorias: treino · equipamento · evento · dúvida · conquista"),
]
y = 1.85
for (t, b) in tables:
    rect(s, 0.6, y, 7.5, 0.86, fill=CARD, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, 0.8, y + 0.09, 2.2, 0.68, [[(t, 12.5, RED, True)]], anchor=MSO_ANCHOR.MIDDLE)
    text(s, 3.0, y + 0.09, 4.95, 0.68, [[(b, 10.8, BODY, False)]], anchor=MSO_ANCHOR.MIDDLE)
    y += 0.97
stats = [("20", "políticas de Row Level Security ativas"),
         ("8", "índices para consultas de catálogo"),
         ("2", "funções SQL: is_admin() e handle_new_user()"),
         ("314", "linhas de migração versionada (supabase/migrations/)")]
y = 1.85
for (n, d) in stats:
    rect(s, 8.5, y, 4.23, 0.86, fill=RED_SOFT, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, 8.7, y + 0.12, 1.1, 0.62, [[(n, 22, RED, True)]], anchor=MSO_ANCHOR.MIDDLE)
    text(s, 9.85, y + 0.12, 2.75, 0.62, [[(d, 10.8, BODY, False)]], anchor=MSO_ANCHOR.MIDDLE)
    y += 0.97
text(s, 0.6, 6.72, 12.1, 0.35,
     [[("is_admin(): função SECURITY DEFINER que protege o admin mestre; visitantes têm apenas leitura de dados públicos.", 11.5, MUTED, False)]])
footer(s)

# ============================================================= S9 — funcionalidades
s = slide()
header(s, "Resultados funcionais", "Módulos implementados")
mods = [
    ("Início", "Hero, destaques e vitrines curadas"),
    ("Marketplace", "Filtros: tipo, categoria, esporte, condição, loja verificada"),
    ("Lojas Oficiais", "Criação de loja e selo “Verificado” via CNPJ"),
    ("Vender C2C", "Anúncio de seminovos com custódia anti-golpe"),
    ("Checkout", "PIX (5% off) com QR Code, cartão e boleto; escrow"),
    ("Treinadores", "Perfis com CREF, avaliação e valor/hora"),
    ("Comunidade", "Posts, curtidas e comentários por categoria"),
    ("Painel Admin", "Auditoria de lojas, pedidos e custódia (is_admin)"),
]
x0, y0 = 0.6, 1.85
cw, ch, gx, gy = 6.0, 1.02, 0.13, 0.13
for i, (t, b) in enumerate(mods):
    col, row = i % 2, i // 2
    x = x0 + col * (cw + gx)
    y = y0 + row * (ch + gy)
    rect(s, x, y, cw, ch, fill=CARD, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(s, x + 0.18, y + 0.18, 0.2, 0.2, fill=RED)
    text(s, x + 0.5, y + 0.14, 2.3, 0.7, [[(t, 13, INK, True)]], anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 2.85, y + 0.14, cw - 3.05, 0.7, [[(b, 11, BODY, False)]], anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.6, 6.55, 12.1, 0.4,
     [[("Nota de transparência: ", 11.5, MUTED, True),
       ("o módulo SmartScan (IA) existe como componente prototipado (SmartScanView.tsx), ainda não integrado à navegação.", 11.5, MUTED, False)]])
footer(s)

# ============================================================= S10 — métricas
s = slide()
header(s, "Resultados quantitativos", "Métricas do código e do banco")
nums = [("86", "arquivos TypeScript (.ts/.tsx)"),
        ("16.641", "linhas de código TypeScript"),
        ("7.769", "linhas em componentes de UI"),
        ("1.973", "linhas na camada de serviços"),
        ("54 / 38", "lojas / produtos no seed de testes"),
        ("20", "políticas RLS no schema")]
x0, y0 = 0.6, 1.85
cw, ch, gx, gy = 3.95, 1.95, 0.16, 0.18
for i, (n, d) in enumerate(nums):
    col, row = i % 3, i // 3
    x = x0 + col * (cw + gx)
    y = y0 + row * (ch + gy)
    rect(s, x, y, cw, ch, fill=CARD, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(s, x + 0.25, y + 0.25, 0.22, 0.22, fill=RED)
    text(s, x + 0.2, y + 0.55, cw - 0.4, 0.75, [[(n, 30, RED, True)]], align=PP_ALIGN.CENTER)
    text(s, x + 0.2, y + 1.32, cw - 0.4, 0.5, [[(d, 11.5, BODY, False)]], align=PP_ALIGN.CENTER)
text(s, 0.6, 6.55, 12.1, 0.4,
     [[("Medições diretas no repositório (wc -l / contagens de schema), outubro de 2026; seed em lib/initial-data.ts.", 11.5, MUTED, False)]])
footer(s)

# ============================================================= S11 — conclusão
s = slide()
header(s, "Conclusão", "Trabalhos futuros e síntese")
text(s, 0.62, 1.8, 12.0, 0.4, [[("TRABALHOS FUTUROS", 11.5, RED, True)]])
bullets(s, 0.6, 2.2, 12.1, 2.6, [
    "Integrar o SmartScan à navegação e conectar a inferência real da API Gemini.",
    "Conectar gateway de pagamento real — hoje o fluxo roda em sandbox com custódia simulada.",
    "Ativar CI com GitHub Actions (.github/workflows/) e testes end-to-end.",
    "Evoluir para experiência mobile (PWA/app) com notificações de pedido.",
], size=14.5, after=9)
card(s, 0.6, 4.75, 12.13, 1.75, "Síntese", [
    "O CoreMotiom entrega um ecossistema esportivo vertical em que confiança é recurso de produto: lojas verificadas via CNPJ, "
    "escrow na revenda C2C e Row Level Security em nível de banco — com governança de código documentada e pronta para evolução.",
], body_size=13.5)
footer(s)

prs.save(OUT)
print("OK", OUT, "| slides:", len(prs.slides.__iter__.__self__._sldIdLst))
